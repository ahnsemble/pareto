import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "artifacts/v3-calibration/dattebayo_v1_9_2.xlsx"
OUT_DIR = ROOT / "app/lib/pareto-store/formula/calibration"
JSON_PATH = OUT_DIR / "dattebayo_coefficients.json"
TS_PATH = OUT_DIR / "dattebayo_coefficients.ts"
SOURCE_URL = "https://docs.google.com/spreadsheets/d/1UQ00geSXjL3lRvrt5G80flwKyuCt5TZY8V_dF8Bw9D0/copy"
EXPORT_URL = "https://docs.google.com/spreadsheets/d/1UQ00geSXjL3lRvrt5G80flwKyuCt5TZY8V_dF8Bw9D0/export?format=xlsx"
SIO_TOOLS_URL = "https://sio-tools.vercel.app/"


NAMED_RANGES = {
    "stage1_base_normalization": [
        "player_atk",
        "Multipliers_Output",
        "Dmg_output",
    ],
    "stage2_equipment_weapon": [
        "Chaos_ssweapon_total",
        "passives_ssweapon",
        "Esuit_dmg",
        "Vneck_dmg",
        "Vgloves_dmg",
        "Vtreads_dmg",
        "AoQ_dmg",
        "TwistingBelt_dmg",
        "SoD_dmg",
        "af_e_ssweapon",
        "af_v_ssweapon",
        "af_c_ssweapon",
        "active_ssweapon",
    ],
    "stage3_tech_parts": [
        "ce_ssweapon",
        "ce_destroyer",
        "ce_drill",
        "ce_molotov",
        "ce_rpg",
        "ce_TBdurian",
        "ce_TBdestroyer",
        "ce_TBforcefield",
        "ce_TBdrill",
        "ce_TBrpg",
        "ce_TBlightning",
    ],
    "stage4_pet_resonance": [
        "XenoResonanceBuff_capy",
        "XenoResonanceBuff_clucker",
        "XenoResonanceBuff_puffo",
        "xenoDamage_capy",
        "xenoDamage_clucker",
        "xenoDamage_puffo",
    ],
    "stage5_collectibles": [
        "critR_relics",
        "critD_relics",
        "skillDmg_synergy",
        "vuln_synergy",
        "shieldDmg_synergy",
    ],
    "stage6_lme_turf": [
        "LMEmultipliersDebuff",
        "LMEdebuff_Turf1",
        "LMEdebuff_Turf2",
        "LMEdebuff_CritR",
        "LMEdebuff_CritD",
        "LMEdebuff_SkillDmg",
        "LMEdebuff_Vuln",
        "LMEdebuff_ShieldDmg",
        "LMEdebuff_toPoisoned",
        "LMEdebuff_toWeakened",
        "LMEdebuff_toChilled",
        "LMEdebuff_DmgDealt",
    ],
    "stage7_conditionals": [
        "critR_total",
        "critD_base",
        "critD_additional",
        "critD_total",
        "debuffDmg_toPoisoned",
        "debuffDmg_toWeakened",
        "debuffDmg_toChilled",
        "debuffDmg_toExposed",
        "laceration_total",
        "TaloxaBeamDmg",
        "JoeyRipsDmg",
        "Overload_total",
        "weak_spot_total",
        "ssgloves_laser",
    ],
    "stage8_xeno_resonance": [
        "xenoPoisonedDmg_capy",
        "xenoPoisonedDmg_clucker",
        "xenoPoisonedDmg_puffo",
        "xenoWeakenedDmg_capy",
        "xenoWeakenedDmg_clucker",
        "xenoWeakenedDmg_puffo",
        "xenoChilledDmg_capy",
        "xenoChilledDmg_clucker",
        "xenoChilledDmg_puffo",
        "XenoResonanceBuff_clucker",
    ],
    "stage9_finalization": [
        "Multipliers_Output",
        "Dmg_output",
        "outputA",
        "outputB",
        "critRateHome",
    ],
}


ROW_BLOCKS = [
    ("stage2_equipment_weapon", "ATK%", "E", "G", "H", 4, 17),
    ("stage5_collectibles", "ATK%", "J", "L", "M", 4, 30),
    ("stage7_conditionals", "CritRate", "E", "G", "H", 4, 18),
    ("stage7_conditionals", "CritRate", "J", "L", "M", 4, 15),
    ("stage7_conditionals", "CritDamage", "E", "G", "H", 9, 16),
    ("stage7_conditionals", "SkillDamage", "E", "G", "H", 4, 18),
    ("stage7_conditionals", "Vulnerability", "E", "G", "H", 4, 18),
    ("stage7_conditionals", "ShieldDamage", "E", "G", "H", 4, 18),
    ("stage7_conditionals", "DebuffDamage", "E", "G", "I", 4, 20),
    ("stage7_conditionals", "UniqueMultipliers", "Q", "T", "U", 4, 12),
    ("stage6_lme_turf", "LMEdebuffCalc", "C", "E", "F", 4, 12),
]


def cell_value(workbook, sheet, coord):
    return workbook[sheet][coord].value


def as_json_value(value):
    if isinstance(value, float):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, bool):
        return value
    if value is None:
        return None
    return str(value)


def defined_cell(wb_formula, wb_values, name, stage):
    destination = list(wb_formula.defined_names[name].destinations)[0]
    sheet, coord = destination
    raw = cell_value(wb_formula, sheet, coord)
    cached = cell_value(wb_values, sheet, coord)
    return {
        "stage": stage,
        "named_range": name,
        "sheet": sheet,
        "cell": coord.replace("$", ""),
        "value": as_json_value(cached),
        "formula": raw if isinstance(raw, str) and raw.startswith("=") else None,
    }


def row_cell(wb_formula, wb_values, stage, sheet, label_col, value_col, gate_col, row):
    label = cell_value(wb_values, sheet, f"{label_col}{row}")
    raw = cell_value(wb_formula, sheet, f"{value_col}{row}")
    cached = cell_value(wb_values, sheet, f"{value_col}{row}")
    gate_raw = cell_value(wb_formula, sheet, f"{gate_col}{row}")
    gate_cached = cell_value(wb_values, sheet, f"{gate_col}{row}")
    return {
        "stage": stage,
        "sheet": sheet,
        "cell": f"{value_col}{row}",
        "label": as_json_value(label),
        "value": as_json_value(cached),
        "formula": raw if isinstance(raw, str) and raw.startswith("=") else None,
        "gate_cell": f"{gate_col}{row}",
        "gate_value": as_json_value(gate_cached),
        "gate_formula": gate_raw if isinstance(gate_raw, str) and gate_raw.startswith("=") else None,
    }


def numeric_named(cells, name):
    for cell in cells:
        if cell.get("named_range") == name:
            value = cell["value"]
            if isinstance(value, (int, float)):
                return float(value)
    raise KeyError(name)


def parse_flat_attack_bonus(cells):
    formula = next(cell["formula"] for cell in cells if cell.get("named_range") == "Multipliers_Output")
    match = re.search(r"player_atk\s*\+\s*([0-9.]+)", formula)
    if not match:
        raise ValueError("flat attack bonus was not found in Multipliers_Output")
    return float(match.group(1))


def build_payload():
    wb_formula = openpyxl.load_workbook(XLSX_PATH, data_only=False, read_only=False)
    wb_values = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=False)

    cells = []
    seen = set()
    for stage, names in NAMED_RANGES.items():
        for name in names:
            if name not in wb_formula.defined_names:
                continue
            entry = defined_cell(wb_formula, wb_values, name, stage)
            key = (entry["stage"], entry.get("named_range"), entry["sheet"], entry["cell"])
            if key not in seen:
                cells.append(entry)
                seen.add(key)

    for stage, sheet, label_col, value_col, gate_col, start, end in ROW_BLOCKS:
        for row in range(start, end + 1):
            entry = row_cell(wb_formula, wb_values, stage, sheet, label_col, value_col, gate_col, row)
            if entry["label"] is None and entry["value"] is None:
                continue
            if sheet == "ATK%" and value_col == "L" and row in (24, 25):
                continue
            key = (entry["stage"], entry["sheet"], entry["cell"], entry.get("label"))
            if key not in seen:
                cells.append(entry)
                seen.add(key)

    flat_attack_bonus = parse_flat_attack_bonus(cells)
    crit_expected = 2 + numeric_named(cells, "critD_additional")
    laceration_factor = 1 + numeric_named(cells, "laceration_total")
    weakened_factor = 1 + numeric_named(cells, "debuffDmg_toWeakened")

    formula_coefficients = {
        "stage1_base_normalization": {
            "baseline_player_atk": numeric_named(cells, "player_atk"),
            "flat_attack_bonus": flat_attack_bonus,
        },
        "stage2_equipment_weapon": {
            "ss_active_slot_bonus": 0.4,
            "ss_af_step_bonus": 0.1,
            "ss_chaos_step_bonus": 0.1,
            "weapon_af_step_bonus": 0.1,
            "ssweapon_chaos_total": numeric_named(cells, "Chaos_ssweapon_total"),
            "ssweapon_passive_factor": numeric_named(cells, "passives_ssweapon"),
        },
        "stage3_tech_parts": {
            "tech_base_gain": 0.05,
            "tech_twinborn_gain": 0.15,
            "tech_resonance_chip_gain": 0.0005,
            "destroyer_contribution_share": numeric_named(cells, "ce_destroyer"),
            "ssweapon_contribution_share": numeric_named(cells, "ce_ssweapon"),
        },
        "stage4_pet_resonance": {
            "deployed_pet_gain": numeric_named(cells, "XenoResonanceBuff_clucker") - 1,
            "assist_pet_gain": 0.05,
            "pet_resonance_atk_per_unit": 0.0001,
        },
        "stage5_collectibles": {
            "collectible_unlock_gain": 0.05,
            "crit_rate_relics": numeric_named(cells, "critR_relics"),
            "crit_damage_relics": numeric_named(cells, "critD_relics"),
        },
        "stage6_lme_turf": {
            "turf1_gain": numeric_named(cells, "LMEdebuff_Turf1"),
            "turf2_gain": numeric_named(cells, "LMEdebuff_Turf2"),
            "lme_damage_dealt_delta": numeric_named(cells, "LMEdebuff_DmgDealt"),
        },
        "stage7_conditionals": {
            "venato_hp_scaling_factor": 0.6,
            "taloxa_laceration_factor": laceration_factor,
            "judgment_weakened_factor": weakened_factor,
            "king_crit_expected_factor": crit_expected,
            "lme_boss_phase_weight": 1 + numeric_named(cells, "LMEdebuff_Turf2"),
            "lme_battle_phase_weight": 1 + numeric_named(cells, "LMEdebuff_Turf1"),
            "crit_rate_total": numeric_named(cells, "critR_total"),
            "crit_damage_additional": numeric_named(cells, "critD_additional"),
        },
        "stage8_xeno_resonance": {
            "default_xeno_delta": numeric_named(cells, "XenoResonanceBuff_clucker") - 1,
            "xeno_clucker_resonance_factor": numeric_named(cells, "XenoResonanceBuff_clucker"),
        },
        "stage9_finalization": {
            "mode_lme_weight": 1.0,
            "mode_ee_weight": 1.0,
            "mode_generic_weight": 1.0,
            "dattebayo_multipliers_output": numeric_named(cells, "Multipliers_Output"),
            "dattebayo_damage_output": numeric_named(cells, "Dmg_output"),
        },
    }

    return {
        "metadata": {
            "source_url": SOURCE_URL,
            "export_url": EXPORT_URL,
            "sio_tools_url": SIO_TOOLS_URL,
            "workbook_path": str(XLSX_PATH),
            "workbook_sheet_count": len(wb_formula.sheetnames),
            "parsed_cell_count": len(cells),
            "formula_stage_count": len(formula_coefficients),
        },
        "formula_coefficients": formula_coefficients,
        "cells": cells,
    }


def write_outputs(payload):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    JSON_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    coefficient_literal = json.dumps(payload["formula_coefficients"], ensure_ascii=False, indent=2)
    TS_PATH.write_text(
        "export const DATTEBAYO_FORMULA_COEFFICIENTS = "
        + coefficient_literal
        + " as const;\n\n"
        + "export type DattebayoFormulaCoefficients = typeof DATTEBAYO_FORMULA_COEFFICIENTS;\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    data = build_payload()
    write_outputs(data)
    print(json.dumps(data["metadata"], ensure_ascii=False, indent=2))
